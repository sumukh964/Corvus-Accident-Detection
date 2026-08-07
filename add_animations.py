# CORVUS presentation -- stage 2: inject slide transitions + entrance animations
# using validated <p:animEffect filter="..."> pattern (tested in isolation first).
from pptx import Presentation
from lxml import etree

PATH = r"C:\Users\shafa\Projects\CORVUS\CORVUS_Presentation.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

def P(tag):
    return f"{{{P_NS}}}{tag}"

def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, tag)
    for k, v in attrs.items():
        el.set(k, str(v))
    return el

TRANSITIONS = [
    ("cut", {}),
    ("push", {"dir": "l"}),
    ("wipe", {"dir": "r"}),
    ("split", {"orient": "horz", "dir": "out"}),
    ("wheel", {"spokes": "4"}),
    ("strips", {"dir": "ld"}),
    ("wedge", {}),
    ("cover", {"dir": "l"}),
    ("pull", {"dir": "r"}),
    ("diamond", {}),
    ("blinds", {"dir": "horz"}),
    ("checker", {"dir": "horz"}),
    ("comb", {"dir": "horz"}),
    ("plus", {}),
    ("randomBar", {"dir": "horz"}),
]

def set_transition(sld, idx):
    existing = sld.find(P("transition"))
    if existing is not None:
        sld.remove(existing)
    name, attrs = TRANSITIONS[idx % len(TRANSITIONS)]
    trans = etree.Element(P("transition"))
    trans.set("spd", "fast")
    child = etree.SubElement(trans, P(name))
    for k, v in attrs.items():
        child.set(k, v)
    csld = sld.find(P("cSld"))
    timing = sld.find(P("timing"))
    if timing is not None:
        timing.addprevious(trans)
    else:
        csld.addnext(trans)

_ID_COUNTER = [1000]
def next_id():
    _ID_COUNTER[0] += 1
    return _ID_COUNTER[0]

FILTERS = ["fly(fromBottom)", "fly(fromLeft)", "wipe(right)", "wipe(left)",
           "blinds(horizontal)", "checkerboard(across)"]

def build_effect_par(shape_id, delay_ms, filt):
    par = etree.Element(P("par"))
    cTn = sub(par, P("cTn"), id=next_id(), presetID="2", presetClass="entr",
              presetSubtype="8", fill="hold", nodeType="afterEffect")
    stC = sub(cTn, P("stCondLst"))
    sub(stC, P("cond"), delay=str(delay_ms))
    childLst = sub(cTn, P("childTnLst"))

    setPar = sub(childLst, P("set"))
    setCbhvr = sub(setPar, P("cBhvr"))
    ctn2 = sub(setCbhvr, P("cTn"), id=next_id(), dur="1", fill="hold")
    stc2 = sub(ctn2, P("stCondLst")); sub(stc2, P("cond"), delay="0")
    tgt2 = sub(setCbhvr, P("tgtEl")); sub(tgt2, P("spTgt"), spid=str(shape_id))
    attrNameLst = sub(setCbhvr, P("attrNameLst"))
    an = sub(attrNameLst, P("attrName")); an.text = "style.visibility"
    toEl = sub(setPar, P("to")); sub(toEl, P("strVal"), val="visible")

    animEffect = sub(childLst, P("animEffect"), transition="in", filter=filt)
    cbhvr = sub(animEffect, P("cBhvr"))
    ctn3 = sub(cbhvr, P("cTn"), id=next_id(), dur="400")
    stc3 = sub(ctn3, P("stCondLst")); sub(stc3, P("cond"), delay="0")
    tgt3 = sub(cbhvr, P("tgtEl")); sub(tgt3, P("spTgt"), spid=str(shape_id))
    return par

def build_timing(shape_ids, filt):
    timing = etree.Element(P("timing"))
    tnLst = sub(timing, P("tnLst"))
    parRoot = sub(tnLst, P("par"))
    cTnRoot = sub(parRoot, P("cTn"), id=next_id(), dur="indefinite", restart="never", nodeType="tmRoot")
    childRoot = sub(cTnRoot, P("childTnLst"))
    seq = sub(childRoot, P("seq"), concat="0", nextAc="seek")
    cTnSeq = sub(seq, P("cTn"), id=next_id(), dur="indefinite", nodeType="mainSeq")
    childSeq = sub(cTnSeq, P("childTnLst"))

    delay = 200
    for sid in shape_ids:
        par = build_effect_par(sid, delay, filt)
        childSeq.append(par)
        delay += 260

    prevCond = sub(seq, P("prevCondLst"))
    c1 = sub(prevCond, P("cond"), evt="onPrev", delay="0")
    t1 = sub(c1, P("tgtEl")); sub(t1, P("sldTgt"))
    nextCond = sub(seq, P("nextCondLst"))
    c2 = sub(nextCond, P("cond"), evt="onNext", delay="0")
    t2 = sub(c2, P("tgtEl")); sub(t2, P("sldTgt"))
    return timing

def set_timing(sld, shape_ids, filt):
    existing = sld.find(P("timing"))
    if existing is not None:
        sld.remove(existing)
    if not shape_ids:
        return
    timing = build_timing(shape_ids, filt)
    sld.append(timing)

def get_top_level_shape_ids(slide, max_shapes=8):
    ids = []
    for shp in slide.shapes:
        ids.append(shp.shape_id)
        if len(ids) >= max_shapes:
            break
    return ids

prs = Presentation(PATH)
for idx, slide in enumerate(prs.slides):
    sld = slide.part._element
    set_transition(sld, idx)
    shape_ids = get_top_level_shape_ids(slide, max_shapes=8)
    filt = FILTERS[idx % len(FILTERS)]
    set_timing(sld, shape_ids, filt)

prs.save(PATH)
print("animations + transitions injected across", len(prs.slides._sldIdLst), "slides")
